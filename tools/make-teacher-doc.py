# -*- coding: utf-8 -*-
"""
生成《课中签到 · 教师使用说明》Word 文档。
排版对照参考文档 3.1.2 节的规范：
  A4，上下 2.54cm / 左右 3.17cm
  正文 仿宋 12pt，首行缩进 21pt，行距 1.5，两端对齐
  一/二级标题 黑体 16pt / 15pt，段前段后 13pt
  图注 仿宋 12pt 居中，「图N. 说明」
  提示词用 1 列 2 行 Table Grid 表格承载
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOT = os.path.join(ROOT, "docs", "screenshots")
OUT = os.path.join(ROOT, "docs", "课中签到-教师使用说明.docx")

FS = "仿宋"      # 正文
HT = "黑体"      # 标题


def set_font(run, name, size, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), name)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)


def style_base(doc):
    st = doc.styles["Normal"]
    st.font.name = FS
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FS)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    for name, size in (("Heading 1", 16), ("Heading 2", 15)):
        h = doc.styles[name]
        h.font.name = HT
        h.font.size = Pt(size)
        h.font.bold = False
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), HT)
        h.paragraph_format.line_spacing = 1.73
        h.paragraph_format.space_before = Pt(13)
        h.paragraph_format.space_after = Pt(13)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.first_line_indent = Pt(0)


def page_setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Pt(595.3), Pt(841.9)
    s.top_margin = s.bottom_margin = Pt(72)
    s.left_margin = s.right_margin = Pt(90)
    s.header_distance, s.footer_distance = Pt(42.55), Pt(49.6)


def add_page_number(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for el, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        e = OxmlElement(el)
        if el == "w:fldChar":
            e.set(qn("w:fldCharType"), attr)
        else:
            e.set(qn("xml:space"), "preserve")
            e.text = " PAGE "
        r._element.append(e)
    set_font(r, FS, 10.5)


def h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    set_font(p.add_run(text), HT, 16)
    return p


def h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    set_font(p.add_run(text), HT, 15)
    return p


def body(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Pt(21) if indent else Pt(0)
    set_font(p.add_run(text), FS, 12)
    return p


def step(doc, num, text):
    """第N步：<加粗标签> + 说明"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Pt(21)
    set_font(p.add_run("第%d步：" % num), FS, 12, bold=True)
    set_font(p.add_run(text), FS, 12)
    return p


def figure(doc, filename, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(os.path.join(SHOT, filename), width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.first_line_indent = Pt(0)
    c.paragraph_format.space_after = Pt(9)
    set_font(c.add_run(caption), FS, 12)


def prompt_table(doc, prompt):
    t = doc.add_table(rows=2, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    head = t.cell(0, 0).paragraphs[0]
    head.paragraph_format.first_line_indent = Pt(0)
    set_font(head.add_run("提示词"), FS, 12, bold=True)

    cell = t.cell(1, 0).paragraphs[0]
    cell.paragraph_format.first_line_indent = Pt(0)
    cell.paragraph_format.line_spacing = 1.3
    first = True
    for line in prompt.split("\n"):
        if not first:
            cell = t.cell(1, 0).add_paragraph()
            cell.paragraph_format.first_line_indent = Pt(0)
            cell.paragraph_format.line_spacing = 1.3
        else:
            first = False
        set_font(cell.add_run(line), FS, 10.5)


PROMPT = """帮我在这台 Windows 电脑上装好一个叫「课中签到」的小工具，我是一线教师，不懂命令行。每一步做完告诉我结果；遇到问题先自己排查，实在不行再问我。

1. 环境：运行 node -v，需要 18 及以上。没有就装：winget install OpenJS.NodeJS.LTS；没有 winget 就从 https://nodejs.org 下载 LTS 安装包装上，装完重开终端再验证。确认有 git；没有就用这个 zip 代替 clone：https://github.com/Caspainnn/class-flow/archive/refs/heads/main.zip

2. 放哪儿：选一个短路径，例如 D:\\ClassFlow。不要放 C:\\Program Files（需要管理员权限，程序写不了数据），也不要放 C:\\Windows 或桌面深层目录。

3. 装依赖：cd D:\\ClassFlow\\server，然后 npm install --omit=dev。如果卡在下载，换国内镜像：npm config set registry https://registry.npmmirror.com

4. 先跑起来验证：node server.js，看到 http://127.0.0.1:3000/ 即成功，先留着这个窗口别关。浏览器打开 http://127.0.0.1:3000/，五个页面（/、/compose、/live、/admin、/join）都要能打开；打不开就把终端里的报错发给我。

5. 换成我的课程和名单（这一步需要我配合）：先问我课程叫什么、学期是什么、名单怎么给我。后台在 http://127.0.0.1:3000/admin，课程管理里可以改课程名、粘贴名单（每行「学号 姓名」）或直接上传 xlsx。名单灌完，确认这门课显示的人数和我学生的人数对得上。

6. 做成双击就能用：pip install pyinstaller（没装 Python 就先 winget install Python.Python.3.12）；在 D:\\ClassFlow 下执行 tools\\build.bat，产出 课中签到.exe；再执行 python tools\\make-shortcut.py 生成桌面快捷方式；关掉第 4 步那个终端，双击桌面图标，确认服务能起来、浏览器能打开。

7. 最后告诉我三样：桌面图标叫什么；启动日志里那行局域网地址（形如 http://192.168.1.5:3000），我投影时要用；以及提醒我手机要和电脑连同一个 Wi-Fi、第一次启动要点「允许」防火墙提示。

注意：data 目录里是我的真实名单，不要上传到任何地方，也不要提交到 git。改任何代码前先备份，改完告诉我改了哪几个文件。"""


def main():
    doc = Document()
    page_setup(doc)
    style_base(doc)
    add_page_number(doc)

    # 标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run("课中签到 · 教师使用说明"), HT, 22)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    set_font(p.add_run("让 WorkBuddy 帮你装好，扫一次码完成点名与随堂一道题"), FS, 12)

    # 1
    h1(doc, "1 这个工具能做什么")
    body(doc, "课中签到把「点名」和「随堂一道题」合成一件事：学生在手机上扫投影上的二维码，答一道选择题，"
              "答完就算签到。老师不用再挨个喊名字，出勤和学情一次拿到。")
    body(doc, "投影页左栏是二维码和每 3 秒刷新的四位课堂码，右栏是本节课还没签到的学生名单。签到一个，"
              "对应的名字就淡下去，位置不动，最后还剩哪几个人一眼就能看见，如图1所示。")
    figure(doc, "live.png", 14.6, "图1. 投影页：左栏二维码与课堂码，右栏未签到名单")

    body(doc, "点「结束本次点名」以后，投影直接变成这道题的选项分布，选错的人多在哪一项看得清清楚楚，"
              "正好接着讲评，如图2所示。")
    figure(doc, "live-result.png", 14.6, "图2. 结束后的讲评页：选项分布")

    body(doc, "所有数据都存在你这台电脑上，不上传任何服务器。课后在后台点「导出」，得到一份 Excel，"
              "包含学号、姓名、题目、提交答案、正确答案、对错、状态和签到时间，可以直接交教务处。")

    # 2
    h1(doc, "2 你需要准备什么")
    body(doc, "三样东西，缺一不可：")
    body(doc, "（1）一台 Windows 10 或 11 的电脑，能上网。第一次安装要下载一些依赖，装好以后断网也能用，"
              "但学生扫码时手机要连到这台电脑上。")
    body(doc, "（2）WorkBuddy 客户端。到 workbuddy.cn 下载安装，用你的账号登录。下面所有操作都通过它完成。")
    body(doc, "（3）上课时手机和电脑在同一个 Wi-Fi 下。学生的手机要能打开你电脑上的地址，这是扫码签到能用的前提。")
    body(doc, "另外准备两样材料，装的过程中 WorkBuddy 会问你要：一是你的课程名称，例如「人工智能安全」；"
              "二是学生名单，Excel 表格或者「学号 姓名」这样的文本都行。")

    # 3
    h1(doc, "3 安装：把这段要求交给 WorkBuddy")
    body(doc, "这一步不用自己敲任何命令。WorkBuddy 能在这台电脑上执行操作，你把下面这段发给它，"
              "它会自己检查环境、下载、安装、启动，最后在桌面上放一个图标。整个过程大约十分钟，"
              "中途它会停下来问你课程名和名单，直接回答即可。")
    step(doc, 1, "打开 WorkBuddy，新建一个对话，确认它工作在这台电脑上（能读写文件、能执行命令）。")
    step(doc, 2, "准备提示词。把下面这段整块复制，粘到对话框里发送。")
    prompt_table(doc, PROMPT)
    step(doc, 3, "等它做完。它会在对话里一步步报告进度：装 Node.js、拉取代码、安装依赖、启动服务、"
                 "打包桌面图标。它问你课程名和名单时直接回答，不用管它具体在做什么。")
    step(doc, 4, "验收。它做完会告诉你桌面图标的名字（默认是「课中签到」）。双击这个图标，"
                 "浏览器会自动打开一个选课页面，看到你自己那门课，就说明装好了，如图3所示。")
    figure(doc, "home.png", 11.5, "图3. 双击图标后打开的选课页")
    body(doc, "如果双击后弹出对话框说找不到 Node，多半是解压或安装不完整，把对话框截图发给 WorkBuddy，"
              "让它重新装一遍即可。")

    # 4
    h1(doc, "4 把演示数据换成自己的")
    body(doc, "第一次启动会自带一门演示课「高等数学（上）」和 47 个虚构学生，那是给你试手感的。"
              "正式用之前换成自己的，只做一次。")
    step(doc, 1, "在选课页右上角点「名单管理」，进入后台的「课程管理」，如图4所示。")
    figure(doc, "admin.png", 14.6, "图4. 后台的课程管理与名单导入")
    step(doc, 2, "改课程名，或者新建一门课。学期可以不填。")
    step(doc, 3, "在名单框里粘学生名单，一行一个，「学号 姓名」中间用空格、逗号或 Tab 隔开都可以，"
                 "全角逗号也认。名单在 Excel 里的话，直接点「上传 xlsx / csv」，它会取前两列当学号和姓名，"
                 "第一行是表头会自动跳过。")
    step(doc, 4, "粘完点「导入粘贴的名单」，回到选课页，确认课程卡片上显示的人数和你班上的人数对得上。")

    # 5
    h1(doc, "5 上课当天的操作")
    step(doc, 1, "课前双击桌面图标，选课页会列出最近上过的课，按数字键 1 或者直接点卡片进入。")
    step(doc, 2, "讲到一段落想提问，就在出题页写题：题干打完回车换行，选项一行一个，在最后一行回车会自动多出一个空行，"
                 "最多六个选项；点左边的圆字母标出正确答案，如图5所示。写好按 Ctrl+Enter 开始点名，"
                 "按 Esc 退回选课。")
    figure(doc, "compose.png", 11.5, "图5. 出题页：一个输入框打完题干和选项")
    step(doc, 3, "把投影页拖到投影仪那块屏幕上，按 F11 全屏。学生扫码、填学号姓名、答一道题，"
                 "提交后只会看到「签到成功 · 第 N 位到课」，不会知道自己答对没有，答案只留给你讲评，如图6所示。")
    figure(doc, "student.png", 7.2, "图6. 学生手机上的答题页")
    step(doc, 4, "讲完这道题，点「结束本次点名」，投影换成选项分布，接着讲评。想再出一道题，按 N 回到出题页。")
    step(doc, 5, "课后要留档，「名单管理 → 看板」里每次点名后面都有「导出」，点一下得到 Excel。")

    # 6
    h1(doc, "6 常见问题")
    qa(doc, "学生扫码打不开页面怎么办？",
       "先确认手机和电脑连的是同一个 Wi-Fi。Windows 第一次启动时会弹「是否允许 Node.js 访问网络」，"
       "当时如果点了取消，到「Windows 安全中心 → 防火墙 → 允许应用通过防火墙」里把 Node.js 勾上。"
       "地址以投影页二维码下面那行为准，换了网络就重启一次服务，地址会跟着变。")
    qa(doc, "学生提交后能看到自己的对错吗？",
       "不能。提交后只回「签到成功 · 第 N 位到课」，选项分布只出现在投影讲评页和后台。")
    qa(doc, "不小心把浏览器关了，点名记录还在吗？",
       "在。回到选课页，顶部会出现一条横幅，点一下就回到刚才那场；服务端重启后也会自动接上没结束的会话。")
    qa(doc, "怎么彻底关掉？",
       "双击文件夹里的「停止课中签到.bat」。关掉浏览器不等于关服务，服务还在后台跑着。")
    qa(doc, "这台电脑需要一直开着吗？",
       "需要。服务跑在这台电脑上，关机或休眠，学生就扫不进来了。")
    qa(doc, "换一台电脑怎么办？",
       "把整个文件夹复制过去，双击里面的图标即可。名单和历次记录都在 data 目录里，跟着文件夹一起走。")

    # 结尾
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.first_line_indent = Pt(21)
    set_font(p.add_run("工具是开源的，地址 https://github.com/Caspainnn/class-flow 。"
                       "用下来有哪里别扭，直接告诉我，改起来很快。"), FS, 12)

    doc.save(OUT)
    print("saved:", OUT)
    print("size: %.1f KB" % (os.path.getsize(OUT) / 1024.0))


def qa(doc, q, a):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(q), FS, 12, bold=True)
    body(doc, a)


if __name__ == "__main__":
    main()
